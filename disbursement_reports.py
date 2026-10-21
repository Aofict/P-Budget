"""Export saved monthly plans and actuals without changing project data."""
from decimal import Decimal
from io import BytesIO
import re
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SOURCES = (('domestic', 'เงินกู้ในประเทศ'), ('foreign', 'เงินกู้ต่างประเทศ'), ('revenue', 'เงินรายได้'))


def period(year, mode):
    months = list(range(10, 13)) + list(range(1, 10)) if mode == 'fiscal' else range(1, 13)
    return [f'{year - 1 if mode == "fiscal" and month >= 10 else year}-{month:02d}' for month in months]


def years_for(projects, mode):
    years = set()
    for project in projects:
        details = project.get('details', {})
        timeline = re.search(r'(\d{4})\s*[-–—]\s*(\d{4})', details.get('timeline', ''))
        if timeline:
            start, end = map(int, timeline.groups())
            if 2400 <= start <= end <= 3000:
                years.update(range(start, end + 1))
        for key in (details.get('monthlyDisbursement') or {}).get('months', {}):
            year, month = map(int, key.split('-'))
            years.add(year + (mode == 'fiscal' and month >= 10))
    return sorted(years)


def amount(projects, keys, kind, source):
    values = []
    for project in projects:
        months = (project.get('details', {}).get('monthlyDisbursement') or {}).get('months', {})
        for key in keys:
            entry = months.get(key, {}).get(kind, {})
            if source in entry:
                values.append(Decimal(str(entry[source])))
    return float(sum(values) / 1_000_000) if values else None


def workbook(projects, *, view, mode, year, plan, division, from_year=None, to_year=None, status_label='ทุกสถานะ'):
    years = years_for(projects, mode) if year == 'all' else [int(year)]
    if from_year is not None and to_year is not None:
        years = list(range(from_year, to_year + 1))
    else:
        years = [value for value in years if (from_year is None or value >= from_year) and (to_year is None or value <= to_year)]
    wb = Workbook()
    wb.remove(wb.active)
    month_names = ('ม.ค.', 'ก.พ.', 'มี.ค.', 'เม.ย.', 'พ.ค.', 'มิ.ย.', 'ก.ค.', 'ส.ค.', 'ก.ย.', 'ต.ค.', 'พ.ย.', 'ธ.ค.')
    groups = [('แผนรายปี', [(str(y), period(y, mode)) for y in years])] if view == 'annual' else [
        (f'แผนรายเดือน {y}', [(month_names[int(key[-2:])-1], [key]) for key in period(y, mode)]) for y in years
    ]
    if not groups:
        groups = [('แผนรายปี' if view == 'annual' else 'แผนรายเดือน', [])]
    for name, periods in groups:
        sheet_name = name.rsplit(' ', 1)[-1] if view == 'monthly' and periods else name
        ws = wb.create_sheet(sheet_name)
        columns = 6 + len(periods) + 1
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
        ws.cell(1, 1, f'{name} · {"ปีงบประมาณ" if mode == "fiscal" else "ปีปฏิทิน"} · หน่วย: ล้านบาท')
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
        ws.cell(2, 1, f'{"ชุดมติ ครม." if plan == "cabinet" else "ชุดเปลี่ยนแปลงงบ"} · {division or "ทุกสายงาน"} · {status_label}')
        for column, title in enumerate(('ลำดับ', 'ชื่อย่อ', 'โครงการ', 'สายงาน', 'แผน / ผล', 'แหล่งเงิน'), 1):
            ws.cell(3, column, title)
            if view == 'monthly':
                ws.merge_cells(start_row=3, start_column=column, end_row=4, end_column=column)
        header_end = 4 if view == 'monthly' else 3
        if view == 'monthly' and periods:
            ws.merge_cells(start_row=3, start_column=7, end_row=3, end_column=6+len(periods))
            ws.cell(3, 7, f'{"ปีงบประมาณ" if mode == "fiscal" else "ปี"} {name.rsplit(" ", 1)[-1]}')
        for index, (label, _) in enumerate([*periods, ('รวม', [])]):
            column = 7 + index
            if view == 'monthly' and label == 'รวม':
                ws.cell(3, column, label)
                ws.merge_cells(start_row=3, start_column=column, end_row=4, end_column=column)
            else:
                ws.cell(header_end, column, label)
        keys = [key for _, month_keys in periods for key in month_keys]
        sequence = 0
        subtotal_rows = set()
        for project in projects:
            project_start = ws.max_row + 1
            details = project.get('details', {})
            project_sources = []
            for source, label in SOURCES:
                target, actual = amount([project], keys, plan, source), amount([project], keys, 'actual', source)
                selected = project.get('details', {}).get('fundingSources', '').split(',')
                finances, details = project.get('finances', {}), project.get('details', {})
                funding = {'domestic': finances.get('loan', 0) - (details.get('foreignLoan') or 0), 'foreign': details.get('foreignLoan') or 0, 'revenue': finances.get('revenue', 0)}
                if source not in selected and not funding[source] and target is None and actual is None:
                    continue
                project_sources.append((source, label))
            if not project_sources:
                continue
            for kind, metric in ((plan, 'แผนการเบิกจ่าย'), ('actual', 'ผลการเบิกจ่าย')):
                group_start = ws.max_row + 1
                for source, label in project_sources:
                    values = [amount([project], month_keys, kind, source) for _, month_keys in periods]
                    ws.append([None, None, None, None, metric if ws.max_row+1 == group_start else None, label, *values, amount([project], keys, kind, source)])
                subtotal = []
                for _, month_keys in [*periods, ('รวม', keys)]:
                    recorded = [amount([project], month_keys, kind, source) for source, _ in project_sources]
                    subtotal.append(sum(value for value in recorded if value is not None) if any(value is not None for value in recorded) else None)
                ws.append([None, None, None, None, None, 'รวม', *subtotal])
                subtotal_rows.add(ws.max_row)
                ws.merge_cells(start_row=group_start, start_column=5, end_row=ws.max_row, end_column=5)
            if ws.max_row >= project_start:
                sequence += 1
                for column, value in enumerate((sequence, details.get('shortName', ''), project['name'], details.get('division', '')), 1):
                    ws.cell(project_start, column, value)
                    ws.merge_cells(start_row=project_start, start_column=column, end_row=ws.max_row, end_column=column)
        total_start = ws.max_row + 1
        for kind, metric in ((plan, 'แผนการเบิกจ่าย'), ('actual', 'ผลการเบิกจ่าย')):
            group_start = ws.max_row + 1
            for label, sources in (('เงินกู้', ('domestic', 'foreign')), ('เงินรายได้', ('revenue',)), ('รวม', tuple(source for source, _ in SOURCES))):
                totals = ['รวมทั้งหมด' if ws.max_row+1 == total_start else None, None, None, None, metric if ws.max_row+1 == group_start else None, label]
                for _, month_keys in [*periods, ('รวม', keys)]:
                    recorded = [amount(projects, month_keys, kind, source) for source in sources]
                    totals.append(sum(value for value in recorded if value is not None) if any(value is not None for value in recorded) else None)
                ws.append(totals)
            ws.merge_cells(start_row=group_start, start_column=5, end_row=ws.max_row, end_column=5)
        ws.merge_cells(start_row=total_start, start_column=1, end_row=ws.max_row, end_column=4)
        ws.freeze_panes = f'G{header_end+1}'
        line = Side(style='thin', color='D9DEE5')
        for row in ws:
            for cell in row:
                header = cell.row <= header_end
                total = cell.row >= total_start or cell.row in subtotal_rows and cell.column >= 6
                cell.font = Font(name='Tahoma', size=11, bold=header or total, color='FFFFFF' if header else '243443')
                cell.fill = PatternFill('solid', fgColor='24435D' if header else 'E5EDF4' if total else 'FFFFFF')
                cell.alignment = Alignment(vertical='center', horizontal='left' if cell.row <= 2 else 'center' if header else 'right' if cell.column >= 7 else 'left', wrap_text=cell.column != 3)
                cell.border = Border(top=line, bottom=line, left=line, right=line)
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                if cell.row > header_end and cell.column >= 7:
                    cell.number_format = '#,##0.000;[Red](#,##0.000);0.000'
        for column in range(1, columns+1):
            ws.column_dimensions[get_column_letter(column)].width = {1:9, 2:16, 3:80, 4:15, 5:22, 6:23}.get(column, 17)
        ws.row_dimensions[1].height = 32
        ws.row_dimensions[2].height = 30
        ws.print_title_rows = f'1:{header_end}'
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
