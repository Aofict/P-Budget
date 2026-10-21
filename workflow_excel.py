"""Excel parser adapted from user supplied MeetingBudget_Test/app.py. No app startup or database side effects."""
import os
import re
from datetime import datetime
from io import BytesIO
try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

WORKFLOW_IMPORT_TEMPLATE = {
    'land': ('1.', 'ค่าที่ดิน', 'detail'),
    'construction': ('2.', 'หมวดสิ่งก่อสร้าง', 'group'),
    'pea_section': ('2.1', 'ส่วนที่ กฟภ.ดำเนินการเอง', 'subgroup'),
    '2.1.1': ('2.1.1', '- ค่าวัสดุอุปกรณ์', 'detail'),
    '2.1.2': ('2.1.2', '- ค่าแรงงาน', 'detail'),
    '2.1.3': ('2.1.3', '- ค่าขนส่ง', 'detail'),
    '2.1.4': ('2.1.4', '- ค่าสำรวจออกแบบและควบคุมงาน', 'detail'),
    '2.1.5': ('2.1.5', '- ค่าเบ็ดเตล็ด', 'detail'),
    '2.1.6': ('2.1.6', '- ค่าดำเนินการ', 'detail'),
    '2.1.7': ('2.1.7', '- ค่าสำรองเผื่อปรับราคา', 'detail'),
    '2.1.8': ('2.1.8', '- ค่าเผื่อขาด', 'detail'),
    '2.1.9': ('2.1.9', '- ค่าอื่น ๆ (ระบุ..)', 'detail'),
    '2.1.10': ('2.1.10', '- ค่าใช้จ่ายทางอ้อม', 'detail'),
    'total_pea': ('รวมงาน กฟภ. ดำเนินการเอง', '', 'subtotal'),
    'contract_section': ('2.2', 'ส่วนที่จ้างบริษัทดำเนินการ (ค่าจ้างเหมาก่อสร้าง/ Turnkey)', 'subgroup'),
    '2.2.1': ('2.2.1', 'ค่าจ้างเหมาก่อสร้าง', 'subtotal'),
    '2.2.1.1': ('2.2.1.1', '- งานก่อสร้างสถานีไฟฟ้า', 'detail'),
    '2.2.1.2': ('2.2.1.2', '- งานก่อสร้างระบบจำหน่ายไฟฟ้า', 'detail'),
    '2.2.1.3': ('2.2.1.3', '- งานก่อสร้างสายส่ง', 'detail'),
    '2.2.1.4': ('2.2.1.4', '- งานก่อสร้างระบบผลิตไฟฟ้า', 'detail'),
    '2.2.1.5': ('2.2.1.5', '- งานเพิ่มเบย์', 'detail'),
    '2.2.1.6': ('2.2.1.6', '- งานเพิ่มหม้อแปลง', 'detail'),
    '2.2.1.7': ('2.2.1.7', '- งานจ้างเหมาอื่นๆ (ระบุ..)', 'detail'),
    '2.2.2': ('2.2.2', 'ค่าสำรวจออกแบบและควบคุมงาน', 'detail'),
    '2.2.3': ('2.2.3', 'ค่าเบ็ดเตล็ด', 'detail'),
    '2.2.4': ('2.2.4', 'ค่าดำเนินการ', 'detail'),
    '2.2.5': ('2.2.5', 'ค่าสำรองเผื่อปรับราคา', 'detail'),
    '2.2.6': ('2.2.6', 'ค่าเผื่อขาด', 'detail'),
    '2.2.7': ('2.2.7', 'ค่าอื่น ๆ (ระบุ..)', 'detail'),
    '2.2.8': ('2.2.8', 'ค่าใช้จ่ายทางอ้อม', 'detail'),
    'total_contract': ('รวมงานจ้างเหมา', '', 'subtotal'),
    'total_construction': ('รวมหมวดสิ่งก่อสร้าง', '', 'subtotal'),
    'total_land_construction': ('รวมหมวดที่ดินและสิ่งก่อสร้าง', '', 'subtotal'),
    'direct_cost': ('', 'รวมค่าใช้จ่ายทางตรง (1, 2.1.1-2.1.9, 2.2.1-2.2.7)', 'subtotal'),
    'indirect_cost': ('', 'รวมค่าใช้จ่ายทางอ้อม (2.1.10, 2.2.8)', 'subtotal'),
    'total_expense': ('', 'รวมค่าใช้จ่าย (ทางตรง+ทางอ้อม)', 'subtotal'),
    'interest': ('3.', 'ดอกเบี้ยระหว่างก่อสร้าง', 'detail'),
    'grand_total': ('รวมทั้งสิ้น (ข้อ1+2+3)', '', 'total'),
    'annual_target': ('', 'รวมเป้าหมายการเบิกจ่ายแต่ละปี', 'metric'),
    'annual_investment': ('', 'วงเงินดำเนินการลงทุนในแต่ละปี', 'metric'),
    'carry_in': ('', 'ผูกพันยกมา+ขอตั้งในปี', 'metric'),
    'carry_out': ('', 'ผูกพันยกไป', 'metric'),
}

WORKFLOW_IMPORT_ORDER = list(WORKFLOW_IMPORT_TEMPLATE.keys())
WORKFLOW_DETAIL_PARENT_KEYS = {
    key for key, meta in WORKFLOW_IMPORT_TEMPLATE.items() if meta[2] == 'detail'
}
WORKFLOW_SCAN_ROW_LIMIT = 5000


def _wf_text(value):
    if value is None:
        return ''
    if isinstance(value, bool):
        return '1' if value else '0'
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.strftime('%d/%m/%Y')
    return str(value).strip()


def _wf_code(value):
    s = _wf_text(value).replace('\n', ' ').strip()
    if s.endswith('.0'):
        s = s[:-2]
    # Excel sometimes stores 2.1 as a numeric cell and 1/2/3 as text with a dot.
    aliases = {'1': '1.', '2': '2.', '3': '3.'}
    return aliases.get(s, s)


def _wf_norm(value):
    return re.sub(r'\s+', ' ', _wf_text(value)).strip().lower()


def _wf_num(value):
    if value is None or value == '':
        return 0.0
    if isinstance(value, bool):
        return 1.0 if value else 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(',', '')
    if not text or text == '-':
        return 0.0
    if text.startswith('(') and text.endswith(')'):
        text = '-' + text[1:-1]
    try:
        return float(text)
    except Exception:
        return 0.0


def _wf_year(value):
    m = re.search(r'(25\d{2}|26\d{2})', _wf_text(value))
    return int(m.group(1)) if m else None


def _workflow_key_from_excel(a_value, b_value, e_value=''):
    code = _wf_code(a_value)
    direct_codes = {
        '1.': 'land', '2.': 'construction', '2.1': 'pea_section',
        '2.1.1': '2.1.1', '2.1.2': '2.1.2', '2.1.3': '2.1.3',
        '2.1.4': '2.1.4', '2.1.5': '2.1.5', '2.1.6': '2.1.6',
        '2.1.7': '2.1.7', '2.1.8': '2.1.8', '2.1.9': '2.1.9',
        '2.1.10': '2.1.10', '2.2': 'contract_section', '2.2.1': '2.2.1',
        '2.2.1.1': '2.2.1.1', '2.2.1.2': '2.2.1.2', '2.2.1.3': '2.2.1.3',
        '2.2.1.4': '2.2.1.4', '2.2.1.5': '2.2.1.5', '2.2.1.6': '2.2.1.6',
        '2.2.1.7': '2.2.1.7', '2.2.2': '2.2.2', '2.2.3': '2.2.3',
        '2.2.4': '2.2.4', '2.2.5': '2.2.5', '2.2.6': '2.2.6',
        '2.2.7': '2.2.7', '2.2.8': '2.2.8', '3.': 'interest',
    }
    if code in direct_codes:
        return direct_codes[code]

    a = _wf_norm(a_value)
    b = _wf_norm(b_value)
    e = _wf_norm(e_value)
    text_map_a = {
        _wf_norm('รวมงาน กฟภ. ดำเนินการเอง'): 'total_pea',
        _wf_norm('รวมงานจ้างเหมา'): 'total_contract',
        _wf_norm('รวมหมวดสิ่งก่อสร้าง'): 'total_construction',
        _wf_norm('รวมหมวดที่ดินและสิ่งก่อสร้าง'): 'total_land_construction',
        _wf_norm('รวมทั้งสิ้น (ข้อ1+2+3)'): 'grand_total',
    }
    if a in text_map_a:
        return text_map_a[a]

    # Match the metric/summary rows by stable Thai phrases instead of exact spaces.
    if 'รวมค่าใช้จ่ายทางตรง' in b:
        return 'direct_cost'
    if 'รวมค่าใช้จ่ายทางอ้อม' in b:
        return 'indirect_cost'
    if 'รวมค่าใช้จ่าย (ทางตรง+ทางอ้อม)' in b or 'รวมค่าใช้จ่าย(ทางตรง+ทางอ้อม)' in b:
        return 'total_expense'
    if 'รวมเป้าหมายการเบิกจ่ายแต่ละปี' in b:
        return 'annual_target'
    if 'วงเงินดำเนินการลงทุนในแต่ละปี' in b:
        return 'annual_investment'
    if 'ผูกพันยกมา+ขอตั้งในปี' in e:
        return 'carry_in'
    if 'ผูกพันยกไป' in e:
        return 'carry_out'
    return None


def _workflow_row_payload(ws_values, row_num, key, row_type, future_years,
                          code_override=None, item_override=None, parent_key=None):
    code_value = code_override if code_override is not None else ws_values.cell(row_num, 1).value
    item_value = item_override if item_override is not None else ws_values.cell(row_num, 2).value
    return {
        'key': key,
        'code': _wf_text(code_value),
        'item': _wf_text(item_value),
        'type': row_type,
        'parent_key': parent_key or '',
        'source_row': row_num,
        'quantity': _wf_text(ws_values.cell(row_num, 3).value),
        'unit': _wf_text(ws_values.cell(row_num, 4).value),
        'estimate_total': _wf_num(ws_values.cell(row_num, 5).value),
        'budget_set': _wf_num(ws_values.cell(row_num, 6).value),
        'actual_prev': _wf_num(ws_values.cell(row_num, 7).value),
        'estimate_current': _wf_num(ws_values.cell(row_num, 8).value),
        'target1': [_wf_num(ws_values.cell(row_num, c).value) for c in range(9, 21)],
        'target2': [_wf_num(ws_values.cell(row_num, c).value) for c in range(22, 34)],
        'future': {
            str(year): _wf_num(ws_values.cell(row_num, 35 + idx).value)
            for idx, year in enumerate(future_years[:14])
        },
        # Keep Excel's cached formula results for display until the user edits data.
        'source_total': (
            _wf_num(ws_values.cell(row_num, 49).value)
            if ws_values.cell(row_num, 49).value is not None else None
        ),  # AW
        'source_grand_total': (
            _wf_num(ws_values.cell(row_num, 54).value)
            if ws_values.cell(row_num, 54).value is not None else None
        ),  # BB
        # BA ("สูง-ต่ำกว่างบ") is imported directly from Excel, just like yearly values.
        # BC and BD are deliberately skipped. BE remains the cancellation reason.
        'cost_saving': _wf_num(ws_values.cell(row_num, 50).value),    # AX
        'reserve': _wf_num(ws_values.cell(row_num, 51).value),        # AY
        'cancel_budget': _wf_num(ws_values.cell(row_num, 52).value),  # AZ
        'over_under': _wf_num(ws_values.cell(row_num, 53).value),     # BA
        'cancel_reason': _wf_text(ws_values.cell(row_num, 57).value), # BE
    }



def parse_workflow_excel(file_bytes, filename, meeting_year):
    if load_workbook is None:
        raise RuntimeError('ยังไม่ได้ติดตั้ง openpyxl กรุณารัน build script ใหม่ หรือใช้ pip install openpyxl')
    if not file_bytes:
        raise ValueError('ไฟล์ Excel ว่างเปล่า')
    if len(file_bytes) > 25 * 1024 * 1024:
        raise ValueError('ไฟล์ Excel มีขนาดเกิน 25 MB')

    lower_name = (filename or '').lower()
    if not lower_name.endswith(('.xlsx', '.xlsm')):
        raise ValueError('รองรับไฟล์ .xlsx และ .xlsm เท่านั้น')

    try:
        wb_values = load_workbook(BytesIO(file_bytes), data_only=True, read_only=False, keep_links=False)
        wb_formulas = load_workbook(BytesIO(file_bytes), data_only=False, read_only=False, keep_links=False)
    except Exception as exc:
        raise ValueError(f'เปิดไฟล์ Excel ไม่สำเร็จ: {exc}')

    sheet_name = '016-1' if '016-1' in wb_values.sheetnames else wb_values.sheetnames[0]
    ws = wb_values[sheet_name]
    wsf = wb_formulas[sheet_name]
    warnings = []

    if ws.max_column < 57:
        raise ValueError('โครงสร้าง Excel ไม่ครบถึงคอลัมน์ BE จึงไม่ใช่แบบ Workflow ที่รองรับ')
    if _wf_norm(ws['A7'].value) != _wf_norm('ลำดับ') or _wf_norm(ws['B7'].value) != _wf_norm('รายการ'):
        raise ValueError('ไม่พบหัวตาราง ลำดับ / รายการ ที่แถว 7')
    if 'เหตุผลที่ยกเลิกงบ/งาน' not in _wf_norm(ws['BE7'].value):
        raise ValueError('ไม่พบคอลัมน์ BE เหตุผลที่ยกเลิกงบ/งาน')

    detected = {
        'actual': _wf_year(ws['G9'].value),
        'estimate': _wf_year(ws['H9'].value),
        'target1': _wf_year(ws['I8'].value),
        'target2': _wf_year(ws['V8'].value),
        'future': [_wf_year(ws.cell(8, c).value) for c in range(35, 49)],
    }
    expected = {
        'actual': meeting_year - 1,
        'estimate': meeting_year,
        'target1': meeting_year + 1,
        'target2': meeting_year + 2,
        'future': [meeting_year + 3 + i for i in range(14)],
    }
    future_years = [y if y else expected['future'][i] for i, y in enumerate(detected['future'])]

    year_ok = (
        detected['actual'] == expected['actual'] and
        detected['estimate'] == expected['estimate'] and
        detected['target1'] == expected['target1'] and
        detected['target2'] == expected['target2'] and
        future_years == expected['future']
    )
    if not year_ok:
        warnings.append(
            'ปีใน Excel ไม่ตรงกับรอบประชุมปัจจุบัน: '
            f"คาด {expected['actual']}/{expected['estimate']}/{expected['target1']}/{expected['target2']} "
            f"แต่พบ {detected['actual']}/{detected['estimate']}/{detected['target1']}/{detected['target2']}"
        )

    # Notes start below the data table. This keeps the parser compatible when
    # actual work-detail rows are inserted between a main row and its 'end' marker.
    notes_row = None
    scan_end = min(ws.max_row, WORKFLOW_SCAN_ROW_LIMIT)
    for r in range(10, scan_end + 1):
        if _wf_norm(ws.cell(r, 1).value) == _wf_norm('หมายเหตุ'):
            notes_row = r
            break
    data_end = (notes_row - 1) if notes_row else scan_end
    if ws.max_row > WORKFLOW_SCAN_ROW_LIMIT and not notes_row:
        warnings.append(
            f'อ่านข้อมูลได้ถึงแถว {WORKFLOW_SCAN_ROW_LIMIT}; '
            'ไม่พบแถวหมายเหตุภายในขอบเขตที่รองรับ'
        )

    rows = []
    seen_template = set()
    active_parent = None
    last_template_key = None
    child_count = 0
    ignored_rows = []
    missing_cached = 0

    for r in range(10, data_end + 1):
        a = ws.cell(r, 1).value
        b = ws.cell(r, 2).value
        e = ws.cell(r, 5).value
        b_norm = _wf_norm(b)

        if b_norm == 'start':
            if last_template_key in WORKFLOW_DETAIL_PARENT_KEYS:
                active_parent = last_template_key
            continue
        if b_norm == 'end':
            active_parent = None
            continue

        key = _workflow_key_from_excel(a, b, e)
        if key:
            active_parent = None
            meta = WORKFLOW_IMPORT_TEMPLATE[key]
            row = _workflow_row_payload(
                ws, r, key, meta[2], future_years,
                code_override=meta[0], item_override=meta[1],
            )
            rows.append(row)
            seen_template.add(key)
            last_template_key = key
            if meta[2] == 'detail':
                active_parent = key

            # Warn if an editable numeric source formula has no cached value.
            for c in [5, 6, 7, 8] + list(range(9, 21)) + list(range(22, 34)) + list(range(35, 49)) + [49, 50, 51, 52, 54, 57]:
                formula = wsf.cell(r, c).value
                cached = ws.cell(r, c).value
                if isinstance(formula, str) and formula.startswith('=') and cached is None:
                    missing_cached += 1
            continue

        # A real Excel file may contain one or many work-detail rows between
        # a main category row and the 'end' marker. Preserve them in the app.
        if active_parent:
            meaningful = any(
                ws.cell(r, c).value not in (None, '')
                for c in list(range(1, 9)) + list(range(9, 21)) + list(range(22, 58))
            )
            if meaningful:
                child_count += 1
                child_key = f'child:{active_parent}:{r}:{child_count}'
                child = _workflow_row_payload(
                    ws, r, child_key, 'detail_child', future_years,
                    parent_key=active_parent,
                )
                rows.append(child)
            continue

        # Ignore blank/support rows, but mention unexpected labeled rows in preview.
        if _wf_text(a) or (_wf_text(b) and b_norm not in ('start', 'end')):
            ignored_rows.append(r)

    # Make sure the app always receives the canonical rows even when a workbook
    # omits a blank section. Missing rows are added as empty placeholders.
    by_key = {row['key']: row for row in rows if row.get('key') in WORKFLOW_IMPORT_TEMPLATE}
    children_by_parent = {}
    for row in rows:
        if row.get('type') == 'detail_child':
            children_by_parent.setdefault(row.get('parent_key'), []).append(row)

    canonical = []
    for key in WORKFLOW_IMPORT_ORDER:
        if key in by_key:
            canonical.append(by_key[key])
        else:
            code, item, row_type = WORKFLOW_IMPORT_TEMPLATE[key]
            canonical.append({
                'key': key, 'code': code, 'item': item, 'type': row_type,
                'parent_key': '', 'source_row': None, 'quantity': '', 'unit': '',
                'estimate_total': 0, 'budget_set': 0, 'actual_prev': 0, 'estimate_current': 0,
                'target1': [0] * 12, 'target2': [0] * 12,
                'future': {str(y): 0 for y in expected['future']},
                'source_total': None, 'source_grand_total': None,
                'cost_saving': 0, 'reserve': 0, 'cancel_budget': 0, 'over_under': 0, 'cancel_reason': '',
            })
        canonical.extend(children_by_parent.get(key, []))

    if missing_cached:
        warnings.append(
            f'พบสูตร Excel {missing_cached} ช่องที่ไม่มีค่าคำนวณล่าสุด (cached value); '
            'ช่องดังกล่าวจะนำเข้าเป็น 0 กรุณาเปิดและ Save ไฟล์ด้วย Excel ก่อน Import หากต้องการค่าจากสูตร'
        )
    if ignored_rows:
        warnings.append('พบแถวที่ไม่อยู่ในโครงสร้างมาตรฐานและถูกข้าม: ' + ', '.join(map(str, ignored_rows[:12])))
    if sheet_name != '016-1':
        warnings.append(f'ไม่พบชีต 016-1 จึงใช้ชีตแรก: {sheet_name}')

    project_meta = {
        'project_name': _wf_text(ws['B3'].value),
        'approval_date': _wf_text(ws['B4'].value),
        'timeline': _wf_text(ws['C5'].value),
    }
    workflow = {
        'template_name': 'Book1.xlsx',
        'source_filename': os.path.basename(filename or 'Book1.xlsx'),
        'source_sheet': sheet_name,
        'imported_at': datetime.now().isoformat(timespec='seconds'),
        'preserve_excel_formula_values': True,
        'excel_years': {
            'actual': detected['actual'], 'estimate': detected['estimate'],
            'target1': detected['target1'], 'target2': detected['target2'],
            'future': future_years,
        },
        'rows': canonical,
    }
    preview = {
        'filename': workflow['source_filename'],
        'sheet': sheet_name,
        'project_meta': project_meta,
        'detected_years': workflow['excel_years'],
        'expected_years': expected,
        'year_match': year_ok,
        'can_import': year_ok,
        'imported_rows': len(canonical),
        'detail_children': child_count,
        'template_rows_found': len(seen_template),
        'template_rows_total': len(WORKFLOW_IMPORT_ORDER),
        'warnings': warnings,
    }
    return workflow, preview


