from contextlib import closing
from copy import deepcopy
from pathlib import Path
import io
import json
import sqlite3
import tempfile
import unittest

from openpyxl import Workbook
from server import create_app
from import_meeting_legacy import import_legacy


class WorkspaceTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(dir=Path(__file__).parent/'.test-data')
        self.database=Path(self.temp.name)/'master.db'
        self.app=create_app(self.database);self.app.config['TESTING']=True;self.c=self.app.test_client()
        self.p=self.c.post('/api/projects',json=dict(name='Project "quoted" <safe>',statusId='started',progress=0,
                      finances=dict(loan=6,revenue=4,oh=1,idc=.5),details=dict(director='Director'))).json
        self.m=self.round(1);self.base='/api/meeting-workspaces/'+self.m['id']+'/api'

    def tearDown(self):self.temp.cleanup()

    def round(self,n):
        m=self.c.post('/api/meetings',json=dict(year=2569,round=n,title='Review',date='2026-09-22')).json
        self.c.post('/api/meetings/'+m['id']+'/projects',json={'projectIds':[self.p['id']]})
        return m

    def doc(self):return self.c.get(self.base+'/get-project/'+self.p['code']).json['data']

    def save(self,d):return self.c.post(self.base+'/save-project',json=d)

    def test_round_isolation_carryover_null_and_explicit_zero(self):
        self.round(2)
        d=self.doc();d['budget']={'2569':dict(committed=10,investment_set=3,disbursement_target=4,disbursement_target_post=None,next_year_uncommitted=1),
                                '2570':dict(investment_set=5,investment_change_before=2,investment_change=-1,disbursement_target=3,disbursement_target_post=0,cabinet_target=2)}
        r=self.save(d);self.assertEqual(r.status_code,200,r.json)
        saved=self.doc();self.assertEqual(saved['budget']['2570']['committed'],5)
        self.assertEqual(saved['budget']['2570']['high_target'],-2)
        self.assertEqual(saved['budget']['2571']['committed'],4)
        row=self.c.get(self.base+'/dashboard/2570').json['data'][0]
        self.assertEqual((row['budget_before'],row['budget_after'],row['budget_total_after_change_year']),(5,-1,-1))
        self.assertEqual((row['disbursement_before'],row['disbursement_after']),(3,0))
        self.assertEqual(self.c.get('/api/projects/'+self.p['id']).json,self.p)
        m2=next(m for m in self.c.get('/api/meetings').json['meetings'] if m['round_no']==2)
        self.assertEqual(self.c.get('/api/meeting-workspaces/'+m2['id']+'/api/get-project/'+self.p['code']).json['data']['budget'],{})

    def test_versions_validation_closed_and_snapshot(self):
        d=self.doc();bad=deepcopy(d);bad['budget']={'2579':{}}
        self.assertEqual(self.save(bad).status_code,400)
        bad=deepcopy(d);bad['workflow']={'rows':[dict(key='land',target1=[0]) ]}
        self.assertEqual(self.save(bad).status_code,400)
        bad=deepcopy(d);bad['budget']={'2570':dict(investment_set=float('nan'))}
        self.assertEqual(self.save(bad).status_code,400)
        bad=deepcopy(d);bad['detailBudget']={'changeBudget':[dict(work='x',months=[1],future={})]}
        self.assertEqual(self.save(bad).status_code,400)
        d['meeting_decision']='เห็นชอบ';d['meeting_comment']='Approved';d['loan_total']=123
        self.assertEqual(self.save(d).status_code,200)
        self.assertEqual(self.doc()['loan_total'],6)
        self.assertEqual(self.save(d).status_code,409)
        mb='/api/meetings/'+self.m['id'];m=self.c.get(mb).json
        self.assertEqual(self.c.post(mb+'/close',json={'version':m['version']}).status_code,200)
        self.assertEqual(self.save(self.doc()).status_code,409)
        self.assertEqual(self.c.put(self.base+'/project-master/'+self.p['code'],json={'version':2,'is_active':False}).status_code,409)
        self.assertTrue(self.c.get(self.base+'/config').json['read_only'])
        exported=self.c.get(self.base+'/export');self.assertEqual(exported.status_code,200)
        self.assertEqual(json.loads(exported.data)['projects'][0]['document']['meeting_comment'],'Approved')

    def test_funding_versions_and_metadata_do_not_edit_master(self):
        self.assertEqual(self.c.get(self.base+'/dashboard-fund-sources/2569').json['version'],0)
        payload=dict(version=0,data=dict(pre_budget_loan=1.001,pre_budget_revenue=2))
        self.assertEqual(self.c.post(self.base+'/dashboard-fund-sources/2569',json=payload).status_code,200)
        self.assertEqual(self.c.post(self.base+'/dashboard-fund-sources/2569',json=payload).status_code,409)
        r=self.c.put(self.base+'/project-master/'+self.p['code'],json=dict(version=1,name='overwrite',summary_group='new_project',show_change_summary=True))
        self.assertEqual(r.status_code,200,r.json)
        self.assertEqual(self.doc()['name'],self.p['name'])
        self.assertEqual(self.c.get(self.base+'/project-master').json['data'][0]['summary_group'],'new_project')

    @staticmethod
    def excel(year=2569):
        wb=Workbook();ws=wb.active;ws.title='016-1'
        for key,val in {'A7':'ลำดับ','B7':'รายการ','BE7':'เหตุผลที่ยกเลิกงบ/งาน','G9':year-1,'H9':year,'I8':year+1,'V8':year+2,
                        'A10':'1.','B10':'ค่าที่ดิน','C10':'2','D10':'แปลง','E10':10,'F10':10,'I10':1.234,'V10':2,
                        'BE10':'เหตุผลเดิม'}.items():ws[key]=val
        for i in range(14):ws.cell(8,35+i,year+3+i)
        out=io.BytesIO();wb.save(out);wb.close();return out.getvalue()

    def test_excel_preview_year_check_and_save_workflow(self):
        path=self.base+'/workflow/import-preview'
        r=self.c.post(path,data={'file':(io.BytesIO(self.excel()),'plan.xlsx')})
        self.assertEqual(r.status_code,200,r.json);self.assertTrue(r.json['preview']['can_import'])
        wf=r.json['workflow'];land=next(row for row in wf['rows'] if row['key']=='land')
        self.assertEqual((land['quantity'],land['target1'][0],land['cancel_reason']),('2',1.234,'เหตุผลเดิม'))
        d=self.doc();d['workflow']=wf;self.assertEqual(self.save(d).status_code,200)
        restarted=create_app(self.database).test_client()
        self.assertEqual(restarted.get(self.base+'/get-project/'+self.p['code']).json['data']['workflow'],wf)
        r=self.c.post(path,data={'file':(io.BytesIO(self.excel(2570)),'wrong-year.xlsx')})
        self.assertFalse(r.json['preview']['can_import'])
        self.assertEqual(self.c.post(path,data={'file':(io.BytesIO(b'bad'),'bad.xlsx')}).status_code,400)

    def test_legacy_additive_idempotent_and_full_backup(self):
        source=Path(__file__).parent/'.imports/MeetingBudget_Test/budget.db'
        if not source.exists():self.skipTest('Original user ZIP fixture is not present')
        first=import_legacy(source,self.database);self.assertFalse(first['alreadyImported'])
        second=import_legacy(source,self.database);self.assertTrue(second['alreadyImported'])
        self.assertEqual(len(self.c.get('/api/projects').json['projects']),5)
        mid=first['rounds']['2569'];p=first['projects']['A0003'];base='/api/meeting-workspaces/'+mid+'/api'
        d=self.c.get(base+'/get-project/'+p['code']).json['data']
        self.assertEqual(d['budget']['2571']['disbursement_target_post'],7000)
        self.assertEqual(d['loan_total']+d['revenue_total'],0)
        backup=Path(self.temp.name)/'backup.db';backup.write_bytes(self.c.get('/api/backup/database').data)
        with closing(sqlite3.connect(backup)) as conn:
            self.assertFalse(conn.execute('PRAGMA foreign_key_check').fetchall())
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM meeting_legacy_imports').fetchone()[0],1)

    def test_zero_budget_draft_and_same_origin(self):
        r=self.c.post('/api/projects',json=dict(name='Draft',statusId='started',progress=0,finances={}))
        self.assertEqual(r.status_code,201,r.json);self.assertEqual(r.json['budget'],0)
        d=self.doc()
        self.assertEqual(self.c.post(self.base+'/save-project',json=d,headers={'Origin':'https://other.example'}).status_code,403)
        self.assertEqual(self.c.get('/meeting-app/'+self.m['id']+'/../../server.py').status_code,404)


if __name__=='__main__':unittest.main()
