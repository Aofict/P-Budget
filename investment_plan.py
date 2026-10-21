"""Read annual investment tables without changing project financing totals."""
import math
import re
from io import BytesIO
from zipfile import ZipFile, BadZipFile
from openpyxl import load_workbook


def validate_plan(plan, error):
    if plan is None:
        return None
    if not isinstance(plan, dict) or set(plan) != {'sheet', 'columns', 'rows'}:
        raise error('ข้อมูลแผนการลงทุนไม่ถูกต้อง')
    columns, rows = plan['columns'], plan['rows']
    if not isinstance(plan['sheet'], str) or len(plan['sheet']) > 100:
        raise error('ชื่อชีตไม่ถูกต้อง')
    if not isinstance(columns, list) or not 2 <= len(columns) <= 60 or any(not isinstance(c, str) or len(c) > 500 for c in columns):
        raise error('หัวตารางแผนการลงทุนไม่ถูกต้อง')
    if not isinstance(rows, list) or len(rows) > 500:
        raise error('แผนการลงทุนรองรับไม่เกิน 500 แถว')
    for row in rows:
        if not isinstance(row, list) or len(row) != len(columns):
            raise error('จำนวนคอลัมน์แผนการลงทุนไม่ตรงกัน')
        for cell in row:
            if cell is None:
                continue
            if isinstance(cell, bool) or not isinstance(cell, (str, int, float)) or (isinstance(cell, str) and len(cell) > 2000) or (isinstance(cell, (int, float)) and not math.isfinite(cell)):
                raise error('ค่าในตารางแผนการลงทุนไม่ถูกต้อง')
    return plan


def read_investment_workbook(content, error, sheet_name=None, list_only=False):
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 256 * 1024 * 1024:
                raise error('ข้อมูลภายในไฟล์ Excel เกิน 256 MB กรุณาคัดลอกชีตแผนการลงทุนไปไฟล์ใหม่แล้วนำเข้า')
        values = load_workbook(BytesIO(content), data_only=True, read_only=True)
        formulas = load_workbook(BytesIO(content), data_only=False, read_only=True)
    except (BadZipFile, ValueError, KeyError, OSError):
        raise error('อ่านไฟล์ Excel ไม่ได้ กรุณาใช้ไฟล์ .xlsx')
    plans = []
    try:
        if list_only:
            return values.sheetnames
        if sheet_name is not None and sheet_name not in values.sheetnames:
            raise error('ไม่พบชีตที่เลือกในไฟล์ Excel')
        for sheet in values:
            if sheet_name is not None and sheet.title != sheet_name:
                continue
            # Ignore formatted trailing rows/columns beyond the annual table area.
            data = list(sheet.iter_rows(max_row=min(sheet.max_row, 5000), max_col=min(sheet.max_column, 100), values_only=True))
            header = next((i for i, row in enumerate(data[:20]) if str(row[0] or '').strip() == 'รายการ'), None)
            if header is None:
                continue
            headings = data[header]
            end = next((i for i, cell in enumerate(headings) if str(cell or '').strip().startswith('คงเหลือ')), None)
            if end is None or not any(str(c or '').strip().startswith(('25', 'ปี 25')) for c in headings):
                continue
            # A:B may be merged for the item header, with descriptions in B.
            item_width = 2 if headings[1] is None else 1
            selected = []
            seen = set()
            for j in range(item_width, end + 1):
                heading = str(headings[j] or '').strip()
                compact = re.sub(r'\s+', '', heading)
                if compact.startswith('วงเงินเต็มโครงการหลัก'):
                    kind = 'วงเงินเต็มโครงการหลัก'
                elif compact.startswith('วงเงินเต็มโครงการ'):
                    kind = 'วงเงินเต็มโครงการ'
                elif re.match(r'^(?:ปี)?\d{4}', compact):
                    selected.append((j, heading))
                    continue
                elif compact == 'รวม':
                    kind = 'รวม'
                elif compact.startswith('คงเหลือ'):
                    kind = 'คงเหลือ'
                else:
                    continue
                if kind not in seen:
                    selected.append((j, kind))
                    seen.add(kind)
            columns = ['รายการ'] + [heading for _, heading in selected]
            rows = []
            formula_rows = list(formulas[sheet.title].iter_rows(max_row=len(data), max_col=len(headings), values_only=True))
            for i, row in enumerate(data[header + 1:], header + 1):
                label = ' '.join(str(c).strip() for c in row[:item_width] if c is not None)
                cells = [row[j] for j, _ in selected]
                if not label and all(c is None for c in cells):
                    if rows:
                        break
                    continue
                for (j, _), cell in zip(selected, cells):
                    formula = formula_rows[i][j]
                    if isinstance(formula, str) and formula.startswith('=') and cell is None:
                        raise error(f'ชีต {sheet.title} มีสูตรที่ยังไม่มีผลคำนวณ กรุณาคำนวณและบันทึกใน Excel ก่อนนำเข้า')
                rows.append([label] + cells)
                if label.replace(' ', '') in ('รวมทั้งสิ้น', 'รวมทั้งหมด'):
                    break
            plans.append(validate_plan({'sheet': sheet.title, 'columns': columns, 'rows': rows}, error))
    finally:
        values.close()
        formulas.close()
    if not plans:
        raise error('ไม่พบตารางแรกของแผนการลงทุนที่มีหัวข้อ รายการ ปี และคงเหลือ')
    return plans
