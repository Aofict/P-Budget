import unittest
from io import BytesIO
from openpyxl import Workbook
from investment_plan import read_investment_workbook, validate_plan
from project_profile import profile_details


class InvestmentPlanTests(unittest.TestCase):
    def workbook(self, formula=False):
        workbook=Workbook();sheet=workbook.active;sheet.title='แผน'
        sheet.append(['แผนการลงทุนรายปี'])
        sheet.append(['รายการ',None,'วงเงินเต็มโครงการ',2571,'2571 เพิ่มเติม','รวม','คงเหลือ'])
        sheet.append(['เงินกู้ในประเทศ'])
        sheet.append([None,'ค่าวัสดุ',100,40,10,'=SUM(D4:E4)' if formula else 50,50])
        stream=BytesIO();workbook.save(stream);return stream.getvalue()

    def test_rows_and_columns_preserved(self):
        plan=read_investment_workbook(self.workbook(),ValueError)[0]
        self.assertEqual(plan['columns'],['รายการ','วงเงินเต็มโครงการ','2571','2571 เพิ่มเติม','รวม','คงเหลือ'])
        self.assertEqual(plan['rows'][1],['ค่าวัสดุ',100,40,10,50,50])
        details=profile_details({'details':{'investmentPlan':plan}},None,ValueError)
        self.assertEqual(details['investmentPlan'],plan)
        self.assertEqual(profile_details({}, {'details':details},ValueError)['investmentPlan'],plan)

    def test_uncalculated_formula_rejected(self):
        with self.assertRaisesRegex(ValueError,'ผลคำนวณ'):
            read_investment_workbook(self.workbook(True),ValueError)

    def test_only_first_table_is_imported(self):
        workbook=Workbook();sheet=workbook.active
        sheet.append(['รายการ','วงเงินเต็มโครงการ',2571,'รวม','คงเหลือ','ตารางอื่น'])
        sheet.append(['ค่าวัสดุ',100,100,100,0,999])
        sheet.append(['รวมทั้งสิ้น',100,100,100,0,999])
        sheet.append([None,None,None,1000])
        sheet.append(['หมายเหตุ','ไม่ใช่ตาราง'])
        stream=BytesIO();workbook.save(stream)
        plan=read_investment_workbook(stream.getvalue(),ValueError)[0]
        self.assertEqual(len(plan['rows']),2)
        self.assertEqual(len(plan['columns']),5)

    def test_bad_data_rejected(self):
        with self.assertRaises(ValueError):
            validate_plan({'sheet':'x','columns':['รายการ','ปี'],'rows':[['a',float('nan')]]},ValueError)
        with self.assertRaises(ValueError):
            read_investment_workbook(b'not excel',ValueError)

    def test_requested_columns_only(self):
        workbook=Workbook();sheet=workbook.active
        sheet.append(['รายการ','วงเงินเต็มโครงการหลัก\n75/25','คิดเป็น %','วงเงินเต็มโครงการ\n75/25','คิดเป็น %','วงเงินเต็มโครงการ\n70/30',2569,'2569 เพิ่มเติม','รวม','คิดเป็น %','คงเหลือ'])
        sheet.append(['รวมทั้งสิ้น',100,75,90,70,80,40,50,90,100,0])
        stream=BytesIO();workbook.save(stream)
        plan=read_investment_workbook(stream.getvalue(),ValueError)[0]
        self.assertEqual(plan['columns'],['รายการ','วงเงินเต็มโครงการหลัก','วงเงินเต็มโครงการ','2569','2569 เพิ่มเติม','รวม','คงเหลือ'])
        self.assertEqual(plan['rows'][0],['รวมทั้งสิ้น',100,90,40,50,90,0])

    def test_selected_sheet_only(self):
        workbook=Workbook();sheet=workbook.active;sheet.title='เลือกชีตนี้'
        sheet.append(['รายการ',2571,'รวม','คงเหลือ'])
        sheet.append(['รวมทั้งสิ้น',10,10,0])
        other=workbook.create_sheet('ไม่ต้องนำเข้า')
        other.append(['รายการ',2571,'รวม','คงเหลือ'])
        other.append(['รวมทั้งสิ้น',20,'=SUM(B2)',0])
        stream=BytesIO();workbook.save(stream);content=stream.getvalue()
        self.assertEqual(read_investment_workbook(content,ValueError,list_only=True),['เลือกชีตนี้','ไม่ต้องนำเข้า'])
        plans=read_investment_workbook(content,ValueError,sheet_name='เลือกชีตนี้')
        self.assertEqual(len(plans),1)
        self.assertEqual(plans[0]['rows'][0][1],10)
        with self.assertRaisesRegex(ValueError,'ไม่พบชีต'):
            read_investment_workbook(content,ValueError,sheet_name='ไม่มีชีตนี้')

    def test_formatted_trailing_cells_do_not_block_import(self):
        workbook=Workbook();sheet=workbook.active
        sheet.append(['รายการ',2571,'รวม','คงเหลือ'])
        sheet.append(['รวมทั้งสิ้น',10,10,0])
        sheet.cell(10000,150).number_format='#,##0'
        stream=BytesIO();workbook.save(stream)
        plan=read_investment_workbook(stream.getvalue(),ValueError)[0]
        self.assertEqual(plan['rows'],[['รวมทั้งสิ้น',10,10,0]])
